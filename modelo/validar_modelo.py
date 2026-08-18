#!/usr/bin/env python3
"""
Validacao independente do modelo da planilha.

Replica em Python a matematica que as formulas do Excel fazem, com os mesmos
inputs, e confere as propriedades que precisam valer:

  1. Black-76: paridade put-call e limites (premio >= valor intrinseco, >= 0).
  2. Hedge com futuro: a reducao de amplitude do resultado tem de ser igual a
     razao de hedge efetiva.
  3. Estruturas: piso da put, teto do collar, ordenacao esperada dos payoffs.
  4. Auditoria de referencias: nenhuma formula da planilha aponta para celula
     vazia (erro silencioso que recalculo nao pega).

Rodar: python modelo/validar_modelo.py
"""

import math
import os
import re
import sys

# --------------------------------------------------------- inputs (= Parametros)
VOL_ANUAL = 5000
ESTOQUE = 1500
PRECO_COMPRA = 1780.0      # R$/saca
PRAZO_MESES = 3.0
CARREGO_MES = 22.0         # R$/saca/mes
ICF_LONGO = 368.40         # US$/saca (Dez/26)
ICF_CURTO = 391.60         # US$/saca (Set/26)
CAMBIO_HOJE = 5.20
CAMBIO_VENC = 5.20
BASIS_COMPRA = -30.0       # R$/saca
BASIS_VENDA = 40.0         # R$/saca
RAZAO_HEDGE = 0.80
SACAS_CONTRATO = 100
PCT_MARGEM = 0.0488
COLCHAO = 3.00
CUSTO_CONTRATO = 25.0      # R$/contrato
VOL_IMP = 0.38
PRAZO_ANOS = 0.25
JUROS = 0.0430

VARIACOES = [-0.35, -0.30, -0.25, -0.20, -0.15, -0.10, -0.05, 0.0,
             0.05, 0.10, 0.15, 0.20, 0.25, 0.30, 0.35]

falhas = []


def checa(condicao, descricao, detalhe=""):
    marca = "OK  " if condicao else "FALHA"
    print(f"  [{marca}] {descricao}" + (f"  -> {detalhe}" if detalhe else ""))
    if not condicao:
        falhas.append(descricao)


def N(x):
    """Normal padrao acumulada - equivalente ao NORMSDIST do Excel."""
    return 0.5 * (1.0 + math.erf(x / math.sqrt(2.0)))


def black76(F, K, sigma, T, r):
    """Retorna (call, put) sobre futuro. Mesma formula das abas Opcoes/Estruturas."""
    st = sigma * math.sqrt(T)
    d1 = (math.log(F / K) + 0.5 * sigma ** 2 * T) / st
    d2 = d1 - st
    desc = math.exp(-r * T)
    call = desc * (F * N(d1) - K * N(d2))
    put = desc * (K * N(-d2) - F * N(-d1))
    return call, put


# ============================================================ 1. Black-76
print("\n1. Black-76 (aba Opcoes)")
F = ICF_LONGO
strikes = []
k0 = round(F * 0.80 / 5) * 5
passo = round(F * 0.05 / 5) * 5
for i in range(13):
    strikes.append(k0 + i * passo)

for K in strikes:
    call, put = black76(F, K, VOL_IMP, PRAZO_ANOS, JUROS)
    desc = math.exp(-JUROS * PRAZO_ANOS)
    # paridade put-call para opcao sobre futuro: C - P = e^(-rT) (F - K)
    paridade = abs((call - put) - desc * (F - K))
    if paridade > 1e-8:
        falhas.append(f"paridade put-call quebrada no strike {K}")
    if put < -1e-9 or call < -1e-9:
        falhas.append(f"premio negativo no strike {K}")
    # premio >= valor intrinseco descontado
    if put < desc * max(K - F, 0) - 1e-8:
        falhas.append(f"put abaixo do intrinseco no strike {K}")

checa(not falhas, "paridade put-call, premios nao-negativos e acima do intrinseco",
      f"{len(strikes)} strikes de US$ {strikes[0]:.0f} a US$ {strikes[-1]:.0f}")

atm = min(strikes, key=lambda k: abs(k - F))
call_atm, put_atm = black76(F, atm, VOL_IMP, PRAZO_ANOS, JUROS)
print(f"       strike ATM = US$ {atm:.2f} | put = US$ {put_atm:.2f} "
      f"({put_atm / F:.1%} do futuro) = R$ {put_atm * CAMBIO_VENC:.2f}/saca")
checa(0.02 < put_atm / F < 0.15, "premio da put ATM em faixa plausivel (2% a 15% do futuro)",
      f"{put_atm / F:.2%}")

premios_put = [black76(F, K, VOL_IMP, PRAZO_ANOS, JUROS)[1] for K in strikes]
print(f"       premio medio das puts = US$ {sum(premios_put)/len(premios_put):.2f} "
      f"= R$ {sum(premios_put)/len(premios_put)*CAMBIO_VENC:.2f}/saca")


# ====================================================== 2. Dimensionamento
print("\n2. Dimensionamento")
sacas_alvo = ESTOQUE * RAZAO_HEDGE
contratos = round(sacas_alvo / SACAS_CONTRATO)
sacas_prot = contratos * SACAS_CONTRATO
cobertura = sacas_prot / ESTOQUE
nocional = contratos * SACAS_CONTRATO * ICF_LONGO * CAMBIO_HOJE
margem = nocional * PCT_MARGEM
caixa = margem * COLCHAO
custo_op = contratos * CUSTO_CONTRATO
alta_suportada = (caixa - margem) / (contratos * SACAS_CONTRATO * CAMBIO_HOJE)

print(f"       contratos = {contratos} | sacas protegidas = {sacas_prot} "
      f"| cobertura = {cobertura:.1%}")
print(f"       nocional = R$ {nocional:,.0f} | margem = R$ {margem:,.0f} "
      f"| caixa recomendado = R$ {caixa:,.0f}")
print(f"       alta suportada pelo colchao = US$ {alta_suportada:.2f}/saca "
      f"({alta_suportada/ICF_LONGO:.1%})")

checa(contratos == 12, "contratos calculados batem com o esperado (12)", f"{contratos}")
checa(sacas_prot <= ESTOQUE, "hedge nao excede o estoque fisico (nao virou posicao)",
      f"{sacas_prot} <= {ESTOQUE}")
checa(0.5 <= cobertura <= 1.0, "cobertura dentro da politica (50% a 100%)", f"{cobertura:.1%}")


# ============================================================= 3. Cenarios
print("\n3. Cenarios - fisico x hedge (futuro vendido)")
carrego_total = CARREGO_MES * PRAZO_MESES
res_fisico, res_comb = [], []
for v in VARIACOES:
    icf = ICF_LONGO * (1 + v)
    revenda = icf * CAMBIO_VENC + BASIS_VENDA
    fisico = (revenda - PRECO_COMPRA - carrego_total) * ESTOQUE
    hedge = (ICF_LONGO - icf) * CAMBIO_VENC * SACAS_CONTRATO * contratos - custo_op
    res_fisico.append(fisico)
    res_comb.append(fisico + hedge)

amp_sem = max(res_fisico) - min(res_fisico)
amp_com = max(res_comb) - min(res_comb)
reducao = 1 - amp_com / amp_sem

print(f"       pior sem hedge = R$ {min(res_fisico):,.0f} | pior com hedge = R$ {min(res_comb):,.0f}")
print(f"       melhor sem hedge = R$ {max(res_fisico):,.0f} | melhor com hedge = R$ {max(res_comb):,.0f}")
print(f"       amplitude sem = R$ {amp_sem:,.0f} | com = R$ {amp_com:,.0f} | reducao = {reducao:.1%}")

checa(abs(reducao - cobertura) < 1e-9,
      "reducao de risco do hedge == cobertura efetiva (identidade que tem de valer)",
      f"{reducao:.4%} vs {cobertura:.4%}")
checa(min(res_comb) > min(res_fisico), "hedge eleva o pior caso",
      f"R$ {min(res_comb):,.0f} > R$ {min(res_fisico):,.0f}")
checa(max(res_comb) < max(res_fisico), "hedge reduz o melhor caso (custo real do hedge)",
      f"R$ {max(res_comb):,.0f} < R$ {max(res_fisico):,.0f}")


# =========================================================== 4. Estruturas
print("\n4. Estruturas - sem hedge / futuro / put / collar")
K_put = round(ICF_LONGO * 0.95 / 5) * 5
K_call = round(ICF_LONGO * 1.10 / 5) * 5
_, prem_put = black76(ICF_LONGO, K_put, VOL_IMP, PRAZO_ANOS, JUROS)
prem_call, _ = black76(ICF_LONGO, K_call, VOL_IMP, PRAZO_ANOS, JUROS)
prem_put_brl = prem_put * CAMBIO_VENC
prem_call_brl = prem_call * CAMBIO_VENC
custo_collar = prem_put_brl - prem_call_brl
custo_base = PRECO_COMPRA + carrego_total

print(f"       put {K_put:.0f} = US$ {prem_put:.2f} (R$ {prem_put_brl:.2f}/saca) | "
      f"call {K_call:.0f} = US$ {prem_call:.2f} (R$ {prem_call_brl:.2f}/saca)")
print(f"       custo liquido do collar = R$ {custo_collar:.2f}/saca")

A, B, C, D = [], [], [], []
for v in VARIACOES:
    icf = ICF_LONGO * (1 + v)
    a = icf * CAMBIO_VENC + BASIS_VENDA - custo_base
    b = a + (ICF_LONGO - icf) * CAMBIO_VENC * RAZAO_HEDGE
    c = a + (max(K_put - icf, 0) * CAMBIO_VENC - prem_put_brl) * RAZAO_HEDGE
    d = a + (max(K_put - icf, 0) * CAMBIO_VENC
             - max(icf - K_call, 0) * CAMBIO_VENC - custo_collar) * RAZAO_HEDGE
    A.append(a); B.append(b); C.append(c); D.append(d)

print(f"       {'estrutura':<14}{'pior':>14}{'melhor':>14}{'amplitude':>14}")
for nome, s in (("A) sem hedge", A), ("B) futuro", B), ("C) put", C), ("D) collar", D)):
    print(f"       {nome:<14}{min(s):>14,.2f}{max(s):>14,.2f}{max(s)-min(s):>14,.2f}")

checa(min(C) > min(A), "put eleva o pior caso vs sem hedge")
checa(max(C) < max(A), "put custa premio no melhor caso")
checa(max(C) > max(B), "put preserva mais alta que o futuro vendido",
      f"R$ {max(C):,.2f} > R$ {max(B):,.2f}")
checa(min(D) > min(A), "collar eleva o pior caso")
checa(max(D) <= max(C) + 1e-9, "collar tem teto (nao supera a put seca no melhor caso)")
checa(custo_collar < prem_put_brl, "collar custa menos que a put seca",
      f"R$ {custo_collar:.2f} < R$ {prem_put_brl:.2f}")
checa((max(B) - min(B)) < (max(A) - min(A)), "futuro reduz amplitude")

# ---- teste de viabilidade: a protecao cabe na margem comercial?
print("\n4.1 Teste de viabilidade contra a margem comercial")
margem_comercial = ICF_LONGO * CAMBIO_VENC + BASIS_VENDA - custo_base
print(f"       margem comercial bruta = R$ {margem_comercial:.2f}/saca")
for nome, custo in (("put ATM", put_atm * CAMBIO_VENC),
                    ("put 95%", prem_put_brl),
                    ("collar", custo_collar),
                    ("futuro (custo op.)", custo_op / sacas_prot)):
    pct = custo / margem_comercial
    veredito = ("INVIAVEL" if pct > 1 else "CARO" if pct > 0.5
                else "ACEITAVEL" if pct > 0.25 else "CONFORTAVEL")
    print(f"       {nome:<20} R$ {custo:>8.2f}/saca  = {pct:>7.1%} da margem  -> {veredito}")

checa(margem_comercial > 0, "existe margem comercial a proteger nos precos de exemplo",
      f"R$ {margem_comercial:.2f}/saca")
checa(put_atm * CAMBIO_VENC > margem_comercial,
      "put ATM custa mais que a margem comercial (achado central do estudo)",
      f"R$ {put_atm*CAMBIO_VENC:.2f} > R$ {margem_comercial:.2f}")
checa(custo_collar < 0.5 * margem_comercial,
      "collar cabe em menos de metade da margem comercial",
      f"R$ {custo_collar:.2f} < R$ {0.5*margem_comercial:.2f}")

# colchao de margem: confere a identidade (multiplicador-1) x %margem
print("\n4.2 Colchao de margem")
pct_colchao = (COLCHAO - 1) * PCT_MARGEM
print(f"       colchao {COLCHAO:.1f}x -> aguenta alta de {pct_colchao:.2%} do nocional "
      f"(US$ {alta_suportada:.2f}/saca)")
checa(abs(alta_suportada / ICF_LONGO - pct_colchao) < 1e-9,
      "identidade do colchao: alta suportada == (multiplicador-1) x % de margem",
      f"{alta_suportada/ICF_LONGO:.4%} vs {pct_colchao:.4%}")
checa(alta_suportada / ICF_LONGO > 0.08,
      "colchao aguenta pelo menos 8% de alta (cafe tem vol de ~38% a.a.)",
      f"{alta_suportada/ICF_LONGO:.2%}")

# a curva invertida penaliza o futuro longo - confere o veredito do estudo
carrego_mercado = (ICF_LONGO - ICF_CURTO) * CAMBIO_HOJE
print(f"\n       carrego pago pelo mercado (curto->longo) = R$ {carrego_mercado:.2f}/saca")
print(f"       custo proprio de carrego                 = R$ {-carrego_total:.2f}/saca")
checa(ICF_CURTO > ICF_LONGO, "mercado invertido nos precos de referencia do estudo",
      f"Set US$ {ICF_CURTO:.2f} > Dez US$ {ICF_LONGO:.2f}")
checa(carrego_mercado < 0, "mercado NAO paga o carrego (cobra para carregar)",
      f"R$ {carrego_mercado:.2f}/saca")


# ==================================================== 5. Programa de hedge
print("\n5. Programa de hedge - preco medio ponderado")
tranches = [(300, 368.40, 5.20), (300, 368.40, 5.20),
            (300, 368.40, 5.20), (300, 368.40, 5.20), (0, 368.40, 5.20), (0, 368.40, 5.20)]
tot_sacas = sum(t[0] for t in tranches)
soma = sum(t[0] * (t[1] * t[2] + BASIS_VENDA) for t in tranches)
medio = soma / tot_sacas if tot_sacas else 0
margem_travada = medio - custo_base
print(f"       sacas travadas = {tot_sacas} ({tot_sacas/VOL_ANUAL:.1%} do volume anual)")
print(f"       preco medio travado = R$ {medio:,.2f}/saca | custo base = R$ {custo_base:,.2f}/saca")
print(f"       margem travada = R$ {margem_travada:,.2f}/saca "
      f"(R$ {margem_travada*tot_sacas:,.0f} no total)")
checa(tot_sacas <= VOL_ANUAL, "programa nao trava mais que o volume anual")
checa(abs(medio - (368.40 * 5.20 + BASIS_VENDA)) < 1e-9,
      "media ponderada correta com tranches de mesmo preco")

if margem_travada < 0:
    print("       ATENCAO: com os precos de exemplo o programa trava PREJUIZO de "
          f"R$ {-margem_travada:,.2f}/saca. A planilha sinaliza isso na aba "
          "'Programa de hedge' - e o comportamento esperado, nao um bug.")


# ============================== 6. Auditoria de referencias da planilha
print("\n6. Auditoria de referencias da planilha")
try:
    from openpyxl import load_workbook
except ImportError:
    print("  [PULADO] openpyxl indisponivel")
    load_workbook = None

if load_workbook:
    caminho = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                           "saida", "Estudo_Hedge_Cafe_ICF.xlsx")
    wb = load_workbook(caminho)
    padrao = re.compile(r"(?:'([^']+)'|([A-Za-z][A-Za-z0-9_ ]*))!\$?([A-Z]{1,2})\$?(\d+)")
    n_form = 0
    vazias = []
    proibidas = ["XLOOKUP", "XMATCH", "FILTER(", "UNIQUE(", "SORT(", "SEQUENCE(",
                 "TEXTJOIN", "IFS(", "MAXIFS", "MINIFS", "SWITCH("]
    usa_proibida = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cel in row:
                if isinstance(cel.value, str) and cel.value.startswith("="):
                    n_form += 1
                    for f in proibidas:
                        if f in cel.value.upper():
                            usa_proibida.append(f"{ws.title}!{cel.coordinate}: {f}")
                    for m in padrao.finditer(cel.value):
                        aba = m.group(1) or m.group(2)
                        if aba not in wb.sheetnames:
                            continue
                        alvo = wb[aba][f"{m.group(3)}{m.group(4)}"]
                        if alvo.value is None:
                            vazias.append(f"{ws.title}!{cel.coordinate} -> {aba}!"
                                          f"{m.group(3)}{m.group(4)}")
    print(f"       {n_form} formulas em {len(wb.worksheets)} abas")
    checa(not vazias, "nenhuma formula aponta para celula vazia em outra aba",
          "; ".join(vazias[:5]) if vazias else "")
    checa(not usa_proibida, "nenhuma funcao pos-2007 sem suporte garantido",
          "; ".join(usa_proibida[:5]) if usa_proibida else "")

print("\n" + "=" * 72)
if falhas:
    print(f"RESULTADO: {len(falhas)} FALHA(S)")
    for f in dict.fromkeys(falhas):
        print(f"  - {f}")
    sys.exit(1)
print("RESULTADO: todas as verificacoes passaram.")
